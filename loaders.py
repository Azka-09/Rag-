"""Document validation and loading for BlushAI.

Every supported format (PDF, CSV, XML, DOCX) is converted into a list of
LangChain ``Document`` objects with consistent metadata, so the existing
splitter -> embeddings -> FAISS -> retriever pipeline works unchanged.

Metadata keys used:
    source        original filename
    file_type     "PDF" | "CSV" | "XML" | "DOCX"
    page          0-based page index (PDF only, as produced by PyPDFLoader)
    row_start/row_end        data-row range (CSV)
    section_start/section_end  element / section range (XML, DOCX)
"""
from __future__ import annotations

import codecs
import csv
import io
import os
import tempfile
import zipfile
from dataclasses import dataclass

from langchain_core.documents import Document

SUPPORTED_TYPES = {".pdf": "PDF", ".csv": "CSV", ".xml": "XML", ".docx": "DOCX"}
MAX_FILE_BYTES = 25 * 1024 * 1024  # 25 MB

CSV_ROWS_PER_DOC = 25
XML_CHARS_PER_DOC = 2000
DOCX_CHARS_PER_DOC = 3000


class UserFacingError(Exception):
    """An error whose message is safe and friendly to show to the user."""


class DocumentLoadError(UserFacingError):
    """Raised when a document cannot be validated or parsed."""


@dataclass
class LoadResult:
    documents: list[Document]
    unit: str  # "Pages" | "Rows" | "Sections"
    count: int
    file_type: str


# ----------------------------------------------------------------------------
# Small helpers
# ----------------------------------------------------------------------------
def format_size(num_bytes: int) -> str:
    size = float(num_bytes)
    for unit in ("B", "KB", "MB", "GB"):
        if size < 1024 or unit == "GB":
            return f"{int(size)} {unit}" if unit == "B" else f"{size:.1f} {unit}"
        size /= 1024
    return f"{num_bytes} B"


def location_label(meta: dict) -> str:
    """Human-readable location of a chunk, used in prompts and source lists."""
    page = meta.get("page")
    if isinstance(page, int):
        return f"Page {page + 1}"
    if "row_start" in meta:
        a, b = meta["row_start"], meta.get("row_end", meta["row_start"])
        return f"Row {a}" if a == b else f"Rows {a}\u2013{b}"
    if "section_start" in meta:
        a, b = meta["section_start"], meta.get("section_end", meta["section_start"])
        return f"Section {a}" if a == b else f"Sections {a}\u2013{b}"
    return "Excerpt"


def _decode_text(data: bytes) -> str:
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")  # never fails; best effort for legacy files


# ----------------------------------------------------------------------------
# Validation
# ----------------------------------------------------------------------------
def validate_upload(name: str, data: bytes) -> str:
    """Check extension, size and file signature. Returns the file type label."""
    ext = os.path.splitext(name or "")[1].lower()
    if ext not in SUPPORTED_TYPES:
        raise DocumentLoadError(
            f"\u201c{name}\u201d isn't a supported file type. "
            "Please upload a PDF, CSV, XML or DOCX file."
        )
    if not data:
        raise DocumentLoadError("This file is empty. Please choose a file that has content.")
    if len(data) > MAX_FILE_BYTES:
        raise DocumentLoadError(
            f"This file is {format_size(len(data))}, which is over the "
            f"{format_size(MAX_FILE_BYTES)} limit. Please upload a smaller file."
        )

    head = data[:4096]
    if ext == ".pdf":
        if b"%PDF-" not in head:
            raise DocumentLoadError("This file has a .pdf name but doesn't look like a real PDF.")
    elif ext == ".docx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                if "word/document.xml" not in zf.namelist():
                    raise DocumentLoadError(
                        "This doesn't look like a Word (.docx) document. "
                        "Older .doc files aren't supported."
                    )
        except zipfile.BadZipFile as exc:
            raise DocumentLoadError(
                "This Word file appears to be corrupted or isn't a real .docx file."
            ) from exc
    elif ext == ".xml":
        has_bom = data[:2] in (codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)
        if not has_bom and not head.lstrip(b"\xef\xbb\xbf \t\r\n").startswith(b"<"):
            raise DocumentLoadError("This file doesn't look like XML (it should start with a tag).")
    elif ext == ".csv":
        has_bom = data[:2] in (codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)
        if head.startswith((b"%PDF", b"PK\x03\x04")) or (b"\x00" in head and not has_bom):
            raise DocumentLoadError("This file has a .csv name but contains binary data, not text.")
    return SUPPORTED_TYPES[ext]


# ----------------------------------------------------------------------------
# Format-specific loaders
# ----------------------------------------------------------------------------
def _load_pdf(name: str, data: bytes) -> LoadResult:
    from langchain_community.document_loaders import PyPDFLoader

    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
            tmp.write(data)
            tmp_path = tmp.name
        pages = PyPDFLoader(tmp_path).load()
    except Exception as exc:
        raise DocumentLoadError(
            "We couldn't read this PDF. It may be corrupted or password-protected."
        ) from exc
    finally:
        if tmp_path and os.path.exists(tmp_path):
            os.unlink(tmp_path)

    for page in pages:
        page.metadata["file_type"] = "PDF"
    return LoadResult(pages, "Pages", len(pages), "PDF")


def _load_csv(name: str, data: bytes) -> LoadResult:
    text = _decode_text(data)
    try:
        try:
            dialect = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        csv.field_size_limit(10_000_000)
        rows = [r for r in csv.reader(io.StringIO(text, newline=""), dialect)
                if any(c.strip() for c in r)]
    except csv.Error as exc:
        raise DocumentLoadError(
            "This CSV file is malformed and couldn't be parsed. "
            "Please check for unbalanced quotes or stray characters."
        ) from exc

    if len(rows) < 2:
        raise DocumentLoadError(
            "A CSV needs a header row and at least one row of data. "
            "This file doesn't have both."
        )

    header: list[str] = []
    seen: dict[str, int] = {}
    for i, col in enumerate(rows[0]):
        col = col.strip() or f"column_{i + 1}"
        seen[col] = seen.get(col, 0) + 1
        header.append(col if seen[col] == 1 else f"{col}_{seen[col]}")

    lines: list[tuple[int, str]] = []
    for idx, row in enumerate(rows[1:], start=1):
        names = header + [f"column_{j + 1}" for j in range(len(header), len(row))]
        parts = [f"{names[j]}: {v.strip()}" for j, v in enumerate(row) if v.strip()]
        if parts:
            lines.append((idx, f"Row {idx}: " + "; ".join(parts)))
    if not lines:
        raise DocumentLoadError("This CSV has headers but no readable data rows.")

    docs = []
    for i in range(0, len(lines), CSV_ROWS_PER_DOC):
        batch = lines[i:i + CSV_ROWS_PER_DOC]
        content = "Columns: " + ", ".join(header) + "\n" + "\n".join(l for _, l in batch)
        docs.append(Document(page_content=content, metadata={
            "file_type": "CSV", "row_start": batch[0][0], "row_end": batch[-1][0],
        }))
    return LoadResult(docs, "Rows", len(lines), "CSV")


def _xml_local(tag) -> str:
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def _xml_walk(elem, path: str, out: list[str]) -> None:
    name = _xml_local(elem.tag)
    here = f"{path} > {name}" if path else name
    attrs = " ".join(f'{_xml_local(k)}="{v}"' for k, v in elem.attrib.items())
    text = " ".join((elem.text or "").split())
    if text or attrs:
        out.append(here + (f" [{attrs}]" if attrs else "") + (f": {text}" if text else ""))
    for child in elem:
        _xml_walk(child, here, out)
        tail = " ".join((child.tail or "").split())
        if tail:
            out.append(f"{here}: {tail}")


def _load_xml(name: str, data: bytes) -> LoadResult:
    from defusedxml import ElementTree as SafeET
    from defusedxml.common import DefusedXmlException
    from xml.etree.ElementTree import ParseError

    try:
        root = SafeET.fromstring(data)
    except DefusedXmlException as exc:
        raise DocumentLoadError(
            "This XML uses unsafe constructs (such as external entities) and was blocked."
        ) from exc
    except ParseError as exc:
        raise DocumentLoadError(
            "This XML file is malformed (for example an unclosed or mismatched tag)."
        ) from exc

    root_name = _xml_local(root.tag)
    children = list(root)
    units = children or [root]

    docs: list[Document] = []
    buf: list[str] = []
    start = None
    for i, unit in enumerate(units, start=1):
        lines: list[str] = []
        _xml_walk(unit, root_name if children else "", lines)
        if not lines:
            continue
        if start is None:
            start = i
        buf.extend(lines)
        if sum(len(l) for l in buf) >= XML_CHARS_PER_DOC:
            docs.append(Document(page_content="\n".join(buf), metadata={
                "file_type": "XML", "section_start": start, "section_end": i}))
            buf, start = [], None
    if buf:
        docs.append(Document(page_content="\n".join(buf), metadata={
            "file_type": "XML", "section_start": start, "section_end": len(units)}))
    return LoadResult(docs, "Sections", len(units), "XML")


def _load_docx(name: str, data: bytes) -> LoadResult:
    from docx import Document as DocxDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = DocxDocument(io.BytesIO(data))
    docs: list[Document] = []
    buf: list[str] = []

    def flush() -> None:
        if buf and "".join(buf).strip():
            n = len(docs) + 1
            docs.append(Document(page_content="\n".join(buf), metadata={
                "file_type": "DOCX", "section_start": n, "section_end": n}))
        buf.clear()

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "") if para.style is not None else ""
            if style.startswith("Heading") or style == "Title":
                flush()  # a heading starts a new section
                buf.append(f"# {text}")
            else:
                buf.append(text)
        elif tag == "tbl":
            table = Table(child, doc)
            buf.append("[Table]")
            for row in table.rows:
                cells, last = [], None
                for cell in row.cells:
                    if cell._tc is last:  # merged cells repeat; keep one copy
                        continue
                    last = cell._tc
                    cells.append(" ".join(cell.text.split()))
                if any(cells):
                    buf.append(" | ".join(cells))
        if sum(len(l) for l in buf) >= DOCX_CHARS_PER_DOC:
            flush()
    flush()
    return LoadResult(docs, "Sections", len(docs), "DOCX")


_LOADERS = {"PDF": _load_pdf, "CSV": _load_csv, "XML": _load_xml, "DOCX": _load_docx}


# ----------------------------------------------------------------------------
# Public entry point
# ----------------------------------------------------------------------------
def load_document(name: str, data: bytes) -> LoadResult:
    """Validate and parse an uploaded file into LangChain Documents."""
    file_type = validate_upload(name, data)
    try:
        result = _LOADERS[file_type](name, data)
    except UserFacingError:
        raise
    except Exception as exc:
        raise DocumentLoadError(
            f"We couldn't read this {file_type} file. It may be corrupted or in an unexpected format."
        ) from exc

    result.documents = [d for d in result.documents if d.page_content.strip()]
    if not result.documents:
        hint = (" It may be a scanned document made of images; text recognition (OCR) isn't supported."
                if file_type == "PDF" else "")
        raise DocumentLoadError(f"No readable text was found in this {file_type} file.{hint}")

    for d in result.documents:
        d.metadata["source"] = name
        d.metadata["file_type"] = file_type
    return result
